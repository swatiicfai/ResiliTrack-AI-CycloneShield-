from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()

def add_slide(title, content):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title_placeholder = slide.shapes.title
    body_shape = slide.shapes.placeholders[1]
    
    title_placeholder.text = title
    tf = body_shape.text_frame
    tf.text = content
    return slide

# Title Slide
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "ResiliTrack AI"
subtitle.text = "Anticipatory Disaster Intelligence\nTrack 5: Cyclone Impact & Infrastructure Vulnerability Forecaster"

# Slides
add_slide("The Problem (The Blind Spot)", 
          "• Broad Warnings Fail Local Needs: Weather cones don't indicate specific hyper-local impacts.\n"
          "• Cascading Infrastructure Failure: Disaster commanders struggle to identify cut-off emergency shelters before flooding.\n"
          "• Communication Bottleneck: Emergency officers receive raw satellite data without AI synthesis.")

add_slide("The Solution - ResiliTrack AI", 
          "• What it is: AI-powered predictive risk platform.\n"
          "• Simulate: Hydro-Elevation mapping using Google Earth Engine.\n"
          "• Analyze: Critical Infrastructure Exposure Graph.\n"
          "• Act: Gemini Multimodal Operations Intelligence.\n"
          "• Fund: Pre-Landfall Parametric Insurance Triggers.")

add_slide("The Gemini Advantage", 
          "• The Brain: We use Gemini 1.5 Flash as a Virtual Disaster Commander.\n"
          "• Targeted Actions: Outputs tactical commands instantly.\n"
          "• Inclusive Alerts: Multi-lingual citizen alerts in Bengali, Odia, Hindi, English.\n"
          "• Edge-Case Handling: System scales response proportionally to actual threat telemetry.")

add_slide("Community Impact", 
          "• Saves Lives: Hyper-localized evacuation instead of generic panic.\n"
          "• Protects the Vulnerable: Ensures hospitals are not cut off unexpectedly.\n"
          "• Inclusive Communication: Overcomes language barriers instantly.\n"
          "• Financial Resilience: Parametric insurance triggers immediate payouts before disaster hits.")

prs.save("ResiliTrack_AI_Pitch_Deck.pptx")
print("Presentation saved as ResiliTrack_AI_Pitch_Deck.pptx")
